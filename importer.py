"""Workbook migration with rich-value and conventional drawing photo support."""
import io,json,secrets,zipfile,posixpath,base64
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from PIL import Image,ImageOps
from flask import request,jsonify,g

def extract(raw):
    z=zipfile.ZipFile(io.BytesIO(raw))
    if sum(f.file_size for f in z.infolist())>250*1024*1024: raise ValueError('Workbook expands beyond the 250 MB import limit.')
    w=load_workbook(io.BytesIO(raw),data_only=True)
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    def xml(path): return ET.fromstring(z.read(path)) if path in z.namelist() else None
    rels=xml('xl/_rels/workbook.xml.rels'); targets={r.attrib['Id']:posixpath.normpath(posixpath.join('xl',r.attrib['Target'].lstrip('/'))) if not r.attrib['Target'].startswith('/') else r.attrib['Target'].lstrip('/') for r in rels}
    sheetpaths={s.attrib['name']:targets[s.attrib['{'+ns['r']+'}id']] for s in xml('xl/workbook.xml').find('s:sheets',ns)}
    photo_index=[]; metadata=[]
    richrels=xml('xl/richData/_rels/richValueRel.xml.rels'); rich=xml('xl/richData/rdrichvalue.xml'); link=xml('xl/richData/richValueRel.xml'); meta=xml('xl/metadata.xml')
    if richrels is not None and rich is not None and link is not None and meta is not None:
        image_targets={r.attrib['Id']:posixpath.normpath(posixpath.join('xl/richData',r.attrib['Target'])) for r in richrels}
        links=[image_targets.get(r.attrib.get('{'+ns['r']+'}id')) for r in link]
        for rv in rich:
            vals=[v.text for v in rv if v.tag.split('}')[-1]=='v']
            try: photo_index.append(links[int(vals[0])])
            except (IndexError,ValueError,TypeError): photo_index.append(None)
        # vm is a one-based valueMetadata record. bk/rc@v indexes the XLRICHVALUE futureMetadata blocks.
        future=next((el for el in meta if el.tag.split('}')[-1]=='futureMetadata' and el.attrib.get('name')=='XLRICHVALUE'),None)
        future_indexes=[]
        if future is not None:
            for bk in future:
                rvb=next((el for el in bk.iter() if el.tag.split('}')[-1]=='rvb'),None)
                future_indexes.append(int(rvb.attrib['i']) if rvb is not None else None)
        vm=meta.find('s:valueMetadata',ns)
        if vm is not None:
            for bk in vm:
                rc=bk.find('s:rc',ns)
                try: metadata.append(future_indexes[int(rc.attrib['v'])])
                except (IndexError,TypeError,KeyError): metadata.append(None)
    result=[]
    for sheet in w:
        normalized=sheet.title.strip().lower()
        if normalized not in ('cafe report','service print'): continue
        section='Cafe' if normalized=='cafe report' else 'Service'
        header=None; namecol=None; qtycol=None; ratecol=None
        for row in sheet.iter_rows(min_row=1,max_row=min(20,sheet.max_row)):
            for cell in row:
                label=str(cell.value or '').strip().lower()
                if label in ('item','items','item name','particulars'): header=cell.row; namecol=cell.column
        if not namecol: raise ValueError(f'Could not find an ITEM column in {sheet.title}.')
        for row in sheet.iter_rows(min_row=header,max_row=min(header+2,sheet.max_row)):
            for cell in row:
                label=str(cell.value or '').strip().lower()
                if label in ('actual','actual stock','actual quantity'): qtycol=cell.column
                if label in ('rate','unit rate','unit price'): ratecol=cell.column
        cellphotos={}
        for cell in xml(sheetpaths[sheet.title]).findall('.//s:c',ns):
            if 'vm' in cell.attrib:
                try:
                    ri=metadata[int(cell.attrib['vm'])-1]; path=photo_index[ri]
                    if path: cellphotos[int(''.join(filter(str.isdigit,cell.attrib['r'])))]=z.read(path)
                except (IndexError,KeyError,TypeError): pass
        for im in sheet._images:
            try: cellphotos.setdefault(im.anchor._from.row+1,im._data())
            except (AttributeError,OSError): pass
        for row in range(header+1,sheet.max_row+1):
            name=sheet.cell(row,namecol).value
            if not isinstance(name,str) or not name.strip() or name.strip().lower() in ('item','items','total','grand total'): continue
            name=name.strip(); photo=None
            if row in cellphotos:
                im=ImageOps.exif_transpose(Image.open(io.BytesIO(cellphotos[row]))).convert('RGB'); im.thumbnail((1600,1600)); b=io.BytesIO();im.save(b,'WEBP',quality=72); photo=base64.b64encode(b.getvalue()).decode()
            qty=sheet.cell(row,qtycol).value if qtycol else None; rate=sheet.cell(row,ratecol).value if ratecol else None
            result.append(dict(name=name,section=section,qty=qty if isinstance(qty,(float,int)) and qty>=0 else None,rate=rate if isinstance(rate,(float,int)) and rate>=0 else None,photo=photo))
    if not result: raise ValueError('No item rows found. Expected Cafe report and/or SERVICE PRINT sheets.')
    return result

def register(a):
    app=a.app
    @app.post('/api/import/preview')
    @a.need('MASTER')
    def preview():
        f=request.files.get('file')
        if not f or not f.filename.lower().endswith('.xlsx'): a.fail('Choose an .xlsx workbook.')
        try: items=extract(f.read())
        except Exception as e: a.fail('Workbook could not be imported: '+str(e))
        token=secrets.token_hex(24); temp=a.DATA/'imports'; temp.mkdir(exist_ok=True)
        (temp/(token+'.json')).write_text(json.dumps({'actor':g.user['id'],'items':items}))
        for i in items:
            sid=3 if i['section']=='Cafe' else 2
            i['existing']=bool(a.one('SELECT id FROM items WHERE section_id=? AND name=?',(sid,i['name'])))
            i['has_photo']=bool(i.pop('photo'))
        return jsonify(token=token,items=items)
    @app.post('/api/import/commit')
    @a.need('MASTER')
    def commit():
        import re,uuid
        d=request.json; token=d.get('token','')
        if not re.fullmatch('[a-f0-9]{48}',token): a.fail('Invalid import reference.')
        path=a.DATA/'imports'/(token+'.json')
        if not path.exists(): a.fail('Import preview expired or was already imported.')
        payload=json.loads(path.read_text())
        if payload['actor']!=g.user['id']: a.fail('Import belongs to another account.',403)
        items=payload['items']; values=d.get('values',[])
        if len(items)!=len(values): a.fail('Review all import rows.')
        a.db().execute('BEGIN IMMEDIATE'); m=a.date.today().isoformat()[:7]
        if a.one('SELECT id FROM audit WHERE action=? AND detail=?',('WORKBOOK_IMPORTED',json.dumps({'token':token,'rows':len(items)}))): a.fail('This workbook preview was already imported.',409)
        for r,v in zip(items,values):
            sid=3 if r['section']=='Cafe' else 2; a.editable(sid,m)
            qty=a.number(v['qty']) if v.get('qty') not in ('',None) else 0; rate=a.money(v.get('rate'))
            old=a.one('SELECT * FROM items WHERE section_id=? AND name=?',(sid,r['name']))
            if old and not old['active']: a.fail(f'{r["name"]} is archived. Resolve this item before importing.')
            if old and qty and a.one('SELECT id FROM movements WHERE item_id=?',(old['id'],)): a.fail(f'{r["name"]} already has stock history. Leave its opening quantity blank.')
            photo=a.save_photo(r['photo']) if r.get('photo') else (old['photo'] if old else None)
            if old:
                iid=old['id']; a.db().execute('UPDATE items SET photo=?,rate=? WHERE id=?',(photo,rate if rate is not None else old['rate'],iid))
            else:
                iid=a.db().execute('INSERT INTO items(code,name,section_id,rate,photo,created) VALUES(?,?,?,?,?,?)',('INV-'+uuid.uuid4().hex[:8].upper(),r['name'],sid,rate,photo,a.now())).lastrowid
            if qty: a.db().execute('INSERT INTO movements(item_id,section_id,type,qty,rate,name,note,date,actor,request_id) VALUES(?,?,?,?,?,?,?,?,?,?)',(iid,sid,'OPENING',qty,rate,r['name'],'Workbook opening import',a.date.today().isoformat(),g.user['id'],uuid.uuid4().hex))
        a.audit('WORKBOOK_IMPORTED',{'token':token,'rows':len(items)});a.db().commit();path.unlink();return jsonify(ok=True)
