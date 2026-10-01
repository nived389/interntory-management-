import io
import cloud_files
from pathlib import Path
from xml.sax.saxutils import escape

def make_export(data,kind,fmt,title,root):
    if kind=='inventory':
        headers=['SI.NO','ITEM','PICTURE','PREV MONTH STOCK','NEW STOCK','DAMAGE','ACTUAL','Section','Serial No','Property','Specification','Expected','Difference','Rate INR','Purchase INR','Breakage INR','Status']
        values=[[n,r['name'],'',r['previous'],r['added'],r['damage'],r['actual'],r.get('section',''),r.get('code',''),r.get('property',''),r['specification'],r['expected'],r['difference'],None if r['rate'] is None else r['rate']/100,r['purchase_value']/100,r['breakage_value']/100,r.get('status','CLOSED')] for n,r in enumerate(data,1)]
        piccol=3
    elif kind=='yearly':
        headers=['Month','Closed item counts','Purchase INR','Breakage quantity','Breakage INR','Variance']
        values=[[r['month'],r['items'],r['purchase_value']/100,r['breakage_qty'],r['breakage_value']/100,r['difference']] for r in data]; piccol=None
    elif kind=='assets':
        headers=['SI.NO','Serial No','Asset Name','Location','Picture','Quantity','Rate INR','Total INR','Property','Notes']
        values=[[n,r.get('code',''),r['name'],r.get('location',''),'',r['qty'],None if r.get('rate') is None else r['rate']/100,None if r.get('rate') is None else (r['qty']*(r['rate'] or 0))/100,r.get('property',''),r.get('notes','')] for n,r in enumerate(data,1)]; piccol=5
    else:
        headers=['SI.NO','Serial No','Item','Section','Quantity','Reason','Photo','Rate INR','Value INR','Staff','Date','Note','Username','User ID']
        values=[[n,r.get('code',''),r['name'],r.get('section',''),abs(r['qty']),r.get('reason',''),'',None if r['rate'] is None else r['rate']/100,None if r['rate'] is None else abs(r['qty'])*r['rate']/100,r.get('staff',''),r['date'],r.get('note',''),r.get('username',''),r.get('submitter_id',r.get('actor',''))] for n,r in enumerate(data,1)]; piccol=7
    if kind=='inventory' and any(r.get('status') not in (None,'CLOSED') for r in data): title+=' [LIVE PREVIEW - NOT CLOSED]'
    for row in data:
        if row.get('photo'): cloud_files.local(root,'photos/thumb-'+row['photo'])
    out=io.BytesIO()
    if fmt=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.drawing.image import Image
        from openpyxl.utils import get_column_letter
        w=Workbook(); s=w.active; s.title=kind.title(); s.append([title]); s.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
        if kind=='inventory':
            s.append(['SI.NO','ITEM','PICTURE','MONTH CLOSING']); s.merge_cells('D2:G2'); s.append(headers)
        else: s.append(headers)
        head=s.max_row
        for n,row in enumerate(values):
            # Prevent spreadsheet formula injection from names and notes.
            s.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in row])
            if piccol and data[n].get('photo'):
                path=root/'photos'/('thumb-'+data[n]['photo'])
                if path.exists():
                    from PIL import Image as PI
                    photo=io.BytesIO(); PI.open(path).save(photo,'PNG'); photo.seek(0)
                    im=Image(photo); im.width=48; im.height=48; s.add_image(im,f'{get_column_letter(piccol)}{s.max_row}'); s.row_dimensions[s.max_row].height=40
        for row in s.iter_rows():
            for cell in row:
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if cell.row<=head: cell.fill=PatternFill('solid',fgColor='173F35'); cell.font=Font(color='FFFFFF',bold=True)
        for col in range(1,len(headers)+1): s.column_dimensions[get_column_letter(col)].width=30 if col==2 else 19
        s.sheet_properties.pageSetUpPr.fitToPage=True; s.page_setup.orientation='landscape'; s.page_setup.paperSize=s.PAPERSIZE_A4; s.page_setup.fitToWidth=1; s.page_setup.fitToHeight=0; s.print_title_rows=f'1:{head}'; s.freeze_panes=f'D{head+1}'; s.auto_filter.ref=f'A{head}:{get_column_letter(len(headers))}{s.max_row}'; w.save(out)
    else:
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4,landscape
        styles=getSampleStyleSheet(); styles['Normal'].fontSize=8; styles['Normal'].leading=10; doc=SimpleDocTemplate(out,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
        # Keep operational columns readable. Excel carries the full extended inventory layout.
        indices=[0,8,1,7,2,3,4,5,11,6,12,16] if kind=='inventory' else list(range(6)) if kind=='yearly' else list(range(10)) if kind=='assets' else [0,1,2,3,4,5,6,8,12,13,10,11]
        pdf_labels={0:'No.',8:'Serial No.',1:'Item',7:'Section',2:'Photo',3:'Previous',4:'Added',5:'Breakage',6:'Counted',11:'Expected',12:'Diff',16:'Status'} if kind=='inventory' else {0:'No.',1:'Serial No.',2:'Asset Name',3:'Location',4:'Photo',5:'Qty',6:'Rate INR',7:'Total INR',8:'Property',9:'Notes'} if kind=='assets' else {0:'No.',1:'Serial No.',2:'Item',3:'Section',4:'Qty',5:'Reason',6:'Photo',8:'Value INR',12:'Username',13:'Staff',10:'Date',11:'Note'}
        table=[[Paragraph(escape(pdf_labels.get(i,headers[i])),styles['Normal']) for i in indices]]
        for n,row in enumerate(values):
            cells=[]
            for i in indices:
                if piccol and i==piccol-1 and data[n].get('photo') and (root/'photos'/('thumb-'+data[n]['photo'])).exists(): cells.append(Image(str(root/'photos'/('thumb-'+data[n]['photo'])),width=38,height=38,kind='proportional'))
                else: cells.append(Paragraph(escape(str(row[i] if row[i] is not None else 'Not recorded')),styles['Normal']))
            table.append(cells)
        widths=([24,65,130,60,40,48,44,48,48,46,44,60] if kind=='inventory' else [24,65,125,75,40,35,55,60,75,150] if kind=='assets' else [24,60,95,55,30,90,40,50,70,35,55,85] if kind!='yearly' else None)
        t=Table(table,colWidths=widths,repeatRows=1,hAlign='LEFT'); t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dcece4')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f6f8f6')]),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('BOTTOMPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,0),1,colors.HexColor('#173f35'))]))
        story=[Paragraph(escape('Travellers Cavern + Travelicious'),styles['Title']),Paragraph(escape(title),styles['Heading2']),Spacer(1,14),t,Spacer(1,22),Paragraph('Reviewed by: ____________________    Approved by: ____________________',styles['Normal'])]
        def footer(canvas,doc):
            canvas.setFont('Helvetica',8);canvas.drawRightString(818,12,f'Page {doc.page}')
        doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()
